"""Synthetic, non-commercial source fixtures for the private editor smoke test."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter

output = Path(__file__).resolve().parents[2] / "tmp" / "work-original-fixtures"
output.mkdir(parents=True, exist_ok=True)
document = Document()
section = document.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(0.8)
for name in ["Normal", "Title", "Heading 1"]:
    document.styles[name].font.color.rgb = RGBColor(0, 0, 0)
for border in document.styles.element.xpath(".//w:pBdr"):
    border.getparent().remove(border)
document.styles["Normal"].font.size = Pt(11)
document.add_paragraph("OfferPSP document editor test", "Title")
document.add_paragraph("This is a synthetic test document. It is not a contract, offer or customer instruction.")
document.add_heading("Original text", 1)
document.add_paragraph("Original clause A must remain unchanged in the stored DOCX. Working edits belong to a separate sheet.")
document.add_paragraph("Проверка русского текста. Исходный файл не заменяется рабочими правками.")
document.add_heading("Unconfirmed terms", 1)
document.add_paragraph("No rates, counterparties, contact details or commercial promises are included.")
document.save(output / "OfferPSP_QA_original.docx")
pdf = canvas.Canvas(str(output / "OfferPSP_QA_original.pdf"), pagesize=letter)
pdf.setTitle("OfferPSP PDF editor test")
pdf.setFont("Helvetica-Bold", 18)
pdf.drawString(58, 728, "OfferPSP PDF editor test")
pdf.setFont("Helvetica", 11)
for index, line in enumerate([
    "This is a synthetic test PDF, not a signed agreement or payment offer.",
    "Original PDF clause B must remain unchanged in private storage.",
    "The working sheet can contain edits without replacing this original.",
    "No real customer, price, contact or commercial promise is included.",
]):
    pdf.drawString(58, 682 - index * 24, line)
pdf.showPage()
pdf.save()
print(output)
